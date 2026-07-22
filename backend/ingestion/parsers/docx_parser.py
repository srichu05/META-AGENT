"""python-docx parser that preserves paragraph headings where available."""

from pathlib import Path

from ..cleaning import clean_text
from ..exceptions import DocumentParsingError
from ..types import ParsedBlock, ParsedDocument
from .base import DocumentParser


class DOCXDocumentParser(DocumentParser):
    def parse(self, path: Path) -> ParsedDocument:
        try:
            from docx import Document
        except ImportError as error:
            raise DocumentParsingError("python-docx is not installed.") from error

        try:
            source = Document(path)
        except Exception as error:
            raise DocumentParsingError(f"Unable to open DOCX document: {error}") from error

        blocks = []
        section_heading = None
        for paragraph in source.paragraphs:
            text = clean_text(paragraph.text)
            if not text:
                continue
            style_name = (paragraph.style.name if paragraph.style else "").lower()
            if style_name.startswith("heading"):
                section_heading = text
                continue
            blocks.append(ParsedBlock(text=text, section_heading=section_heading))

        if not blocks:
            raise DocumentParsingError("The DOCX document did not contain extractable text.")
        return ParsedDocument(text="\n\n".join(block.text for block in blocks), blocks=blocks)
